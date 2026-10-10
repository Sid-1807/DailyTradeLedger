import os
import logging
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger("TradeAgent")

COLUMNS = [
    "Trade Date",
    "Order No",
    "Trade No",
    "Exchange",
    "Symbol",
    "Action",
    "Quantity",
    "Gross Rate (₹)",
    "Brokerage (₹)",
    "Net Rate (₹)",
    "Total Amount (₹)",
    "Contract Note File"
]

def get_excel_filepath(action: str, q_string: str, fy_string: str) -> str:
    """Returns absolute file path for Buy or Sell workbook for the given quarter and financial year."""
    prefix = "Buy_Trades" if action.upper() == "BUY" else "Sell_Trades"
    filename = f"{prefix}_{q_string}_{fy_string}.xlsx"
    return os.path.join(config.TRADES_DIR, filename)

def clean_str(val) -> str:
    """Helper to convert cell values to clean string representations."""
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s

def load_existing_trade_keys(file_path: str) -> set:
    """Returns set of unique trade keys present in existing Excel file to avoid duplicates."""
    if not os.path.exists(file_path):
        return set()

    try:
        df = pd.read_excel(file_path, skiprows=2, dtype=str)
        if df.empty or "Order No" not in df.columns:
            return set()

        keys = set()
        for _, row in df.iterrows():
            trade_date = clean_str(row.get("Trade Date"))[:10] # Standardize to YYYY-MM-DD
            if trade_date == "TOTAL":
                continue
            order_no = clean_str(row.get("Order No"))
            trade_no = clean_str(row.get("Trade No"))
            symbol = clean_str(row.get("Symbol"))
            action = clean_str(row.get("Action"))
            if order_no and trade_no:
                keys.add((trade_date, order_no, trade_no, symbol, action))
        return keys
    except Exception as e:
        logger.warning(f"Error reading existing keys from {file_path}: {e}")
        return set()

def style_excel_workbook(file_path: str, action: str, q_string: str, fy_string: str):
    """Applies beautiful styling, currency formatting, auto column widths, and total row."""
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    is_buy = action.upper() == "BUY"
    header_fill_color = "D9E1F2" if is_buy else "FCE4D6" # Light blue for Buy, Soft Coral for Sell
    title_fill_color = "2F5597" if is_buy else "C65911"
    
    # Fonts
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    data_font = Font(name="Calibri", size=11)
    total_font = Font(name="Calibri", size=11, bold=True)

    # Fills & Borders
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    title_fill = PatternFill(start_color=title_fill_color, end_color=title_fill_color, fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    # Apply Title Row (Row 1)
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"{'BUY' if is_buy else 'SELL'} TRADES LEDGER - {q_string} {fy_string}"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Apply Table Headers (Row 3)
    ws.row_dimensions[3].height = 24
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Clean up any existing TOTAL rows before calculating new totals
    for r in range(ws.max_row, 3, -1):
        cell_val = str(ws.cell(row=r, column=1).value or "").strip()
        if cell_val == "TOTAL":
            ws.delete_rows(r)

    # Format Data Rows (Row 4 to data_max_row)
    data_max_row = ws.max_row
    if data_max_row >= 4:
        for r in range(4, data_max_row + 1):
            ws.row_dimensions[r].height = 20
            for c in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = data_font
                cell.border = thin_border
                
                col_name = COLUMNS[c - 1]
                
                if col_name == "Trade Date":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    val_str = str(cell.value or "").strip()
                    if val_str:
                        dt = None
                        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%m-%y", "%d/%m/%y"):
                            try:
                                dt = datetime.strptime(val_str.split()[0], fmt.split()[0]).date()
                                break
                            except ValueError:
                                pass
                        if dt:
                            cell.value = dt
                    cell.number_format = "dd-mm-yy"

                elif col_name in ["Order No", "Trade No"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.value = clean_str(cell.value)
                    cell.number_format = "@"

                elif col_name in ["Exchange", "Action"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif col_name == "Symbol":
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif col_name == "Quantity":
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"

                elif "Rate" in col_name or "Brokerage" in col_name or "Amount" in col_name:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "₹#,##0.00"

                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # Add Single Total Row
    total_row_idx = data_max_row + 1
    ws.cell(row=total_row_idx, column=1, value="TOTAL").font = total_font
    ws.cell(row=total_row_idx, column=1).alignment = Alignment(horizontal="center")
    
    # Formula for Total Quantity (Col 7) and Total Amount (Col 11)
    if data_max_row >= 4:
        qty_cell = ws.cell(row=total_row_idx, column=7, value=f"=SUM(G4:G{data_max_row})")
        qty_cell.font = total_font
        qty_cell.number_format = "#,##0"
        
        amt_cell = ws.cell(row=total_row_idx, column=11, value=f"=SUM(K4:K{data_max_row})")
        amt_cell.font = total_font
        amt_cell.number_format = "₹#,##0.00"

    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=total_row_idx, column=c)
        cell.border = double_bottom_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row == 1: # Skip title row
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)

def log_trades_to_excel(trades: list):
    """
    Groups trades by Quarter, Financial Year, and Action (BUY / SELL),
    and appends new trades into Excel files.
    """
    if not trades:
        logger.info("No trades to log.")
        return 0

    added_count = 0

    # Group trades by (Q, FY, Action)
    grouped = {}
    for t in trades:
        try:
            dt = datetime.strptime(t["trade_date"], "%Y-%m-%d")
        except Exception:
            dt = datetime.now()
        
        q_str, fy_str = config.get_quarter_and_fy(dt)
        act = t["action"].upper()
        key = (q_str, fy_str, act)
        
        if key not in grouped:
            grouped[key] = []
        grouped[key].append(t)

    for (q_str, fy_str, action), trade_list in grouped.items():
        file_path = get_excel_filepath(action, q_str, fy_str)
        existing_keys = load_existing_trade_keys(file_path)

        new_trade_rows = []
        for t in trade_list:
            t_date_std = str(t["trade_date"]).strip()[:10]
            trade_key = (
                t_date_std,
                str(t["order_no"]).strip(),
                str(t["trade_no"]).strip(),
                str(t["symbol"]).strip(),
                str(t["action"]).strip()
            )
            if trade_key not in existing_keys:
                new_trade_rows.append({
                    "Trade Date": t["trade_date"],
                    "Order No": t["order_no"],
                    "Trade No": t["trade_no"],
                    "Exchange": t["exchange"],
                    "Symbol": t["symbol"],
                    "Action": t["action"],
                    "Quantity": t["quantity"],
                    "Gross Rate (₹)": t["gross_rate"],
                    "Brokerage (₹)": t["brokerage"],
                    "Net Rate (₹)": t["net_rate"],
                    "Total Amount (₹)": t["total_amount"],
                    "Contract Note File": t["contract_note_file"]
                })
                existing_keys.add(trade_key)

        if not new_trade_rows:
            logger.info(f"All {len(trade_list)} {action} trades for {q_str} {fy_str} already exist in Excel. Skipping.")
            continue

        # Load existing df or create new
        if os.path.exists(file_path):
            try:
                existing_df = pd.read_excel(file_path, skiprows=2)
                # Filter out old total row if present
                if "Trade Date" in existing_df.columns:
                    existing_df = existing_df[existing_df["Trade Date"] != "TOTAL"]
                new_df = pd.DataFrame(new_trade_rows)
                combined_df = pd.concat([existing_df[COLUMNS], new_df[COLUMNS]], ignore_index=True)
            except Exception as e:
                logger.warning(f"Error reading {file_path}, recreating: {e}")
                combined_df = pd.DataFrame(new_trade_rows)[COLUMNS]
        else:
            combined_df = pd.DataFrame(new_trade_rows)[COLUMNS]

        # Sort combined_df chronologically by Trade Date
        combined_df["_sort_date"] = pd.to_datetime(combined_df["Trade Date"], errors="coerce")
        combined_df = combined_df.sort_values(by="_sort_date", ascending=True).drop(columns=["_sort_date"])

        # Write dataframe starting at row 3 (Row 1=Title, Row 2=Empty, Row 3=Header)
        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            combined_df.to_excel(writer, startrow=2, index=False)

        # Apply styling & auto-formatting
        style_excel_workbook(file_path, action, q_str, fy_str)
        added_count += len(new_trade_rows)
        logger.info(f"Successfully added {len(new_trade_rows)} new {action} trades to {file_path}.")

    return added_count
